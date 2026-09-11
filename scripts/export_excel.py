#!/usr/bin/env python3
"""路书 JSON -> Excel 导出。

用法:
    python export_excel.py roadbook.json [-o 路书.xlsx]

字段:日期 | 历史天气与穿衣推荐 | 海拔 | 行程起点终点 | 风景点 | 封路可能性 | 沿线加油站 | 预估耗油
耗油估算: 基础耗油 = 里程 x 油耗/100;海拔修正 +2%/1000m(按终点海拔);
          爬升修正 +0.15L/1000m 累计爬升;高原综合再 x1.08(低温+风阻)。
"""
import argparse
import sys

from roadbook_utils import estimate_fuel, load_roadbook


def safe_excel_value(value):
    """Prevent user-controlled text from becoming an Excel formula."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value

def main():
    p = argparse.ArgumentParser()
    p.add_argument("roadbook", help="路书 JSON 文件")
    p.add_argument("-o", "--output", default=None)
    args = p.parse_args()

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        sys.exit("需要 openpyxl: pip install openpyxl")

    try:
        rb = load_roadbook(args.roadbook)
    except (OSError, ValueError) as exc:
        sys.exit(str(exc))
    rider = rb.get("rider", {})
    consumption = rider.get("consumption_l_per_100km", 3.5)
    out = args.output or "roadbook.xlsx"

    wb = Workbook(); ws = wb.active; ws.title = "逐日路书"
    headers = ["天数", "日期", "历史天气与穿衣推荐", "海拔", "行程起点终点",
               "路面/难度/可信度", "风景点", "封路可能性", "沿线加油站", "顺路餐饮", "备用方案", "预估耗油(L)"]
    widths = [6, 11, 34, 18, 24, 24, 34, 26, 30, 28, 30, 11]
    hfill = PatternFill("solid", fgColor="2E5E4E")
    for c, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, c, h)
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = hfill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = w

    risk_fill = {"高": PatternFill("solid", fgColor="F4CCCC"),
                 "中": PatternFill("solid", fgColor="FCE8CD")}
    total_fuel = 0.0
    r = 2
    for d in rb["days"]:
        if d.get("is_gap_day"):
            vals = [d["day"], d.get("date") or "机动", "机动缓冲日(封路/绕路/高反延误备用)", "-",
                    "视前一天延误情况", "-", "-", "-", "-", "-", d.get("notes", "-"), "-"]
        else:
            spots = "、".join(f"{s['name']}({s.get('type','')}{'·'+s['level'] if s.get('level') else ''})" for s in d.get("scenic_spots", []))
            fuels = "、".join(f"{s['name']}{'(必加满)' if '必加' in s.get('note','') else ''}"
                              for s in d.get("fuel_stops", []))
            m = d.get("meals", {})
            meals = f"午:{m.get('lunch','')};晚:{m.get('dinner','')}" if m else ""
            fuel = estimate_fuel(d, consumption)
            total_fuel += fuel
            vals = [d["day"], d.get("date", ""),
                    f"{d.get('weather_typical','')};穿衣:{d.get('clothing','')}",
                    f"最高{d.get('max_elevation_m','-')}m / 住{d.get('end_elevation_m','-')}m",
                    f"{d.get('start','')} → {d.get('end','')} ({d.get('distance_km','-')}km)",
                    f"{d.get('surface','待核验')} / 难度{d.get('technical_difficulty','-')} / {d.get('route_confidence','unverified')}",
                    spots, d.get("road_closure_risk", ""), fuels, meals,
                    d.get("plan_b", ""), round(fuel, 1)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, safe_excel_value(v))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        risk = str(d.get("road_closure_risk", ""))
        for key, fill in risk_fill.items():
            if risk.startswith(key):
                ws.cell(r, 8).fill = fill
        if d.get("is_gap_day"):
            for c in range(1, 13):
                ws.cell(r, c).fill = PatternFill("solid", fgColor="EDEDED")
        r += 1

    ws.cell(r, 11, "全程预估总耗油").font = Font(bold=True)
    ws.cell(r, 12, round(total_fuel, 1)).font = Font(bold=True)
    ws.freeze_panes = "A2"
    wb.save(out)
    print(f"已导出 {out}: {len(rb['days'])} 天, 预估总耗油 {round(total_fuel,1)}L")

if __name__ == "__main__":
    main()
