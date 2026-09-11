import copy
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

from roadbook_utils import validate_roadbook, wgs84_to_gcj02  # noqa: E402
from export_excel import safe_excel_value  # noqa: E402


def sample_roadbook():
    base_day = {
        "day": 1,
        "date": "2026-09-25",
        "is_gap_day": False,
        "start": "成都",
        "end": "康定",
        "route_level": "detailed",
        "waypoints": [],
        "route_geometry": [[104.07, 30.67], [103.1, 30.2], [101.96, 30.05]],
        "route_source": {"provider": "高德 MCP", "strategy": "避开高速", "checked_at": "2026-09-10"},
        "route_phase": "outbound",
        "start_coords": [104.07, 30.67],
        "end_coords": [101.96, 30.05],
        "distance_km": 260,
        "max_elevation_m": 3000,
        "end_elevation_m": 2560,
        "scenic_spots": [],
        "scenic_routes": [{"name": "测试风景公路", "reason": "测试连续景观"}],
        "fuel_stops": [],
        "alternative_routes": [{"name": "测试备选", "trigger": "主线管制"}],
        "road_closure_risk": "中:测试",
        "weather_typical": "测试",
        "clothing": "测试",
        "schedule": {"am": "出发", "pm": "抵达", "evening": "休息"},
        "meals": {"lunch": "沿途县城", "dinner": "住宿点附近"},
        "crowd_avoidance": {"enabled": True, "plan": "避开午后拥堵"},
        "snow_risk": {"level": "低", "basis": "历史同期测试资料", "checked_at": "2026-09-10", "action": "出发前复核"},
        "weather_forecast": {"status": "available", "summary": "多云", "updated_at": "2026-09-10"},
        "historical_weather": {
            "probability_summary": "近十年同期雨雪日约20%",
            "precipitation_probability_pct": 20,
            "source": "测试气象资料",
            "checked_at": "2026-09-10",
            "sample_description": "近10年同旬30个样本日",
        },
        "lodging_options": [{"name": "测试酒店", "status": "verify", "source": "飞猪", "checked_at": "2026-09-10"}],
        "context_pois": [{"name": "测试垭口", "type": "山峰/垭口", "lat": 30.2, "lng": 103.1}],
        "plan_b": "原地等待并复核",
        "requires_lodging": True,
    }
    gap_one = {
        **copy.deepcopy(base_day),
        "day": 2,
        "date": None,
        "is_gap_day": True,
        "route_level": "rough",
        "route_phase": "gap",
        "distance_km": 0,
        "requires_lodging": True,
    }
    gap_two = {**copy.deepcopy(gap_one), "day": 3, "requires_lodging": False}
    return {
        "title": "测试</script><script>alert(1)</script>",
        "disclaimer": "仅供参考，请量力而行。出发前复核动态信息。",
        "coordinate_system": "gcj02",
        "rider": {
            "motorcycle_model": "测试车型",
            "experience_level": "intermediate",
            "plateau_experience": "some",
            "offroad_experience": "basic",
            "preferred_pace": "normal",
        },
        "service_preflight": {
            "checked_at": "2026-09-10",
            "services": {
                "amap": {"status": "available", "provider": "高德 MCP"},
                "weather": {"status": "available", "provider": "和风天气"},
                "lodging": {"status": "available", "provider": "飞猪"},
            },
        },
        "crowd_avoidance": {"mode": "avoid", "locations": ["测试地点"]},
        "holiday_strategy": {
            "is_holiday_period": True,
            "day1_extension_proposed": True,
            "user_decision": "accept",
            "plan": "首日延长至康定",
        },
        "days": [base_day, gap_one, gap_two],
    }


class CoreTests(unittest.TestCase):
    def test_valid_sample(self):
        self.assertEqual(validate_roadbook(sample_roadbook()), [])

    def test_sample_matches_json_schema_when_validator_available(self):
        try:
            import jsonschema
        except ImportError:
            self.skipTest("jsonschema 未安装")
        schema = json.loads((ROOT / "references" / "roadbook.schema.json").read_text(encoding="utf-8"))
        jsonschema.validate(sample_roadbook(), schema)

    def test_rejects_missing_coordinate_system(self):
        roadbook = sample_roadbook()
        del roadbook["coordinate_system"]
        self.assertTrue(any("coordinate_system" in item for item in validate_roadbook(roadbook)))

    def test_rejects_missing_disclaimer(self):
        roadbook = sample_roadbook()
        del roadbook["disclaimer"]
        self.assertTrue(any("仅供参考，请量力而行" in item for item in validate_roadbook(roadbook)))

    def test_rejects_high_risk_route_without_safety_fields(self):
        roadbook = sample_roadbook()
        roadbook["days"][0]["technical_difficulty"] = 4
        del roadbook["days"][0]["route_geometry"]
        errors = validate_roadbook(roadbook)
        self.assertTrue(any("route_geometry" in item for item in errors))
        self.assertTrue(any("bailout_points" in item for item in errors))

    def test_rejects_detailed_route_without_real_geometry(self):
        roadbook = sample_roadbook()
        del roadbook["days"][0]["route_geometry"]
        self.assertTrue(any("道路轨迹" in item for item in validate_roadbook(roadbook)))

    def test_rejects_snow_high_without_no_go_conditions(self):
        roadbook = sample_roadbook()
        roadbook["days"][0]["snow_risk"]["level"] = "高"
        self.assertTrue(any("no_go_conditions" in item for item in validate_roadbook(roadbook)))

    def test_rejects_missing_rider_interview(self):
        roadbook = sample_roadbook()
        del roadbook["rider"]["plateau_experience"]
        self.assertTrue(any("plateau_experience" in item for item in validate_roadbook(roadbook)))

    def test_rejects_missing_service_preflight(self):
        roadbook = sample_roadbook()
        del roadbook["service_preflight"]["services"]["lodging"]
        self.assertTrue(any("services.lodging" in item for item in validate_roadbook(roadbook)))

    def test_holiday_requires_day_one_extension_proposal(self):
        roadbook = sample_roadbook()
        roadbook["holiday_strategy"]["day1_extension_proposed"] = False
        self.assertTrue(any("首日延长" in item for item in validate_roadbook(roadbook)))

    def test_detailed_day_requires_weather_probability_and_lodging(self):
        roadbook = sample_roadbook()
        del roadbook["days"][0]["historical_weather"]
        del roadbook["days"][0]["lodging_options"]
        errors = validate_roadbook(roadbook)
        self.assertTrue(any("historical_weather" in item for item in errors))
        self.assertTrue(any("lodging_options" in item for item in errors))

    def test_wgs84_conversion_changes_sichuan_coordinate(self):
        lng, lat = wgs84_to_gcj02(104.07, 30.67)
        self.assertNotAlmostEqual(lng, 104.07, places=4)
        self.assertNotAlmostEqual(lat, 30.67, places=4)

    def test_excel_formula_text_is_neutralized(self):
        self.assertEqual(safe_excel_value("=1+1"), "'=1+1")
        self.assertEqual(safe_excel_value("普通文本"), "普通文本")

    def test_budget_counts_all_days_and_only_required_nights(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "roadbook.json"
            source.write_text(json.dumps(sample_roadbook(), ensure_ascii=False), encoding="utf-8")
            result = subprocess.run(
                [sys.executable, str(SCRIPTS / "budget_estimator.py"), str(source), "--tier", "comfort"],
                check=True,
                capture_output=True,
                text=True,
            )
            self.assertIn("2晚", result.stdout)
            self.assertIn("3天", result.stdout)

    def test_excel_contains_disclaimer(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl 未安装")
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "roadbook.json"
            output = Path(temp_dir) / "roadbook.xlsx"
            source.write_text(json.dumps(sample_roadbook(), ensure_ascii=False), encoding="utf-8")
            subprocess.run([sys.executable, str(SCRIPTS / "export_excel.py"), str(source), "-o", str(output)], check=True)
            workbook = load_workbook(output, read_only=True)
            values = [cell.value for row in workbook.active.iter_rows() for cell in row if cell.value]
            self.assertTrue(any("仅供参考，请量力而行" in str(value) for value in values))

    def test_html_neutralizes_script_terminator(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "roadbook.json"
            output = Path(temp_dir) / "roadbook.html"
            source.write_text(json.dumps(sample_roadbook(), ensure_ascii=False), encoding="utf-8")
            subprocess.run(
                [sys.executable, str(SCRIPTS / "export_html.py"), str(source), "-o", str(output)],
                check=True,
                capture_output=True,
                text=True,
            )
            generated = output.read_text(encoding="utf-8")
            self.assertNotIn("</script><script>alert(1)</script>", generated)
            self.assertIn("&lt;/script&gt;", generated)
            self.assertIn("AMapLoader.load", generated)
            self.assertNotIn("leaflet", generated.lower())
            self.assertIn("查看第${d.day}天详情", generated)
            self.assertIn("仅供参考，请量力而行", generated)
            self.assertIn("主路线来自已核验道路轨迹", generated)
            self.assertIn("去程", generated)
            self.assertIn("回程", generated)
            self.assertIn("驻地环线", generated)
            self.assertIn("outbound:'#ec6a38'", generated)
            self.assertIn("return:'#247985'", generated)
            self.assertIn("local:'#c89438'", generated)
            self.assertIn("'#6aaec4'", generated)
            self.assertIn("酒店 / 民宿建议（飞猪优先）", generated)
            self.assertIn("历史同期概率", generated)
            self.assertIn("mapStyle:'amap://styles/normal'", generated)
            self.assertIn("features:['bg','point','road','building']", generated)
            self.assertIn("d.context_pois", generated)
            if shutil.which("node"):
                scripts = re.findall(r"<script>(.*?)</script>", generated, re.DOTALL)
                result = subprocess.run(["node", "--check"], input=scripts[-1], capture_output=True, text=True)
                self.assertEqual(result.returncode, 0, result.stderr)


if __name__ == "__main__":
    unittest.main()
